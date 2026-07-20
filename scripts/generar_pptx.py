import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SCRIPTS_DIR = os.path.dirname(__file__)
SCREENSHOTS_DIR = os.path.join(SCRIPTS_DIR, "screenshots")
GRAFICOS_DIR = os.path.join(SCRIPTS_DIR, "graficos")
OUTPUT_PATH = os.path.join(SCRIPTS_DIR, "presentacion-cliente.pptx")

NAVY = RGBColor(0x1A, 0x36, 0x5D)
GOLD = RGBColor(0xD6, 0x9E, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x2D, 0x37, 0x48)
GRAY = RGBColor(0x71, 0x71, 0x71)
LIGHT_BG = RGBColor(0xF7, 0xFA, 0xFC)

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=DARK, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf

def add_multi_text(slide, left, top, width, height, items, font_size=14, color=DARK):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"  {chr(0x2022)} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return tf

def add_slide_number(slide, num, total):
    add_text_box(slide, Inches(8.8), Inches(7.1), Inches(1.2), Inches(0.4),
                 f"{num}/{total}", font_size=10, color=GRAY, alignment=PP_ALIGN.RIGHT)

def add_image_safe(slide, path, left, top, width, height=None):
    if os.path.exists(path):
        if height:
            slide.shapes.add_picture(path, left, top, width, height)
        else:
            slide.shapes.add_picture(path, left, top, width)

def create_presentation():
    print("Generando presentación PowerPoint...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    TOTAL_SLIDES = 17
    slide_num = 0

    # ============================================================
    # SLIDE 1: PORTADA
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    add_bg(slide, NAVY)
    slide_num += 1

    add_text_box(slide, Inches(1), Inches(1.5), Inches(8), Inches(1),
                 "MAGNA", font_size=54, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(2.5), Inches(8), Inches(0.7),
                 "Ingeniería y Topografía", font_size=26, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(3.8), Inches(8), Inches(1.2),
                 "Informe de Transformación Digital\nModernización, Rediseño y Nuevas Funcionalidades", font_size=18, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(5.5), Inches(8), Inches(0.5),
                 "Julio 2026", font_size=16, color=GOLD, alignment=PP_ALIGN.CENTER)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 2: RESUMEN
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "¿Qué logramos?", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.2), Inches(9), Inches(0.5),
                 "El sitio web de MAGNA fue transformado por completo. Estos son los resultados:", font_size=14, color=GRAY)

    logros = [
        "Velocidad de carga mejorada en más de 30%",
        "Diseño responsive moderno y profesional",
        "Seguridad con SSL, proxy inverso y protección de datos",
        "Sección Quiénes Somos con misión, visión y valores",
        "Formulario de contacto que envía correos automáticos",
        "Blog funcional con búsqueda y diseño mejorado",
        "Tienda online con carrito y pagos PayPal",
        "Política de tratamiento de datos personales",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.9), Inches(9), Inches(4.5), logros, font_size=15, color=DARK)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 3: PAGESPEED BEFORE/AFTER
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Velocidad y Rendimiento", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Resultados PageSpeed Insights - Comparativa Antes vs Después", font_size=14, color=GRAY)

    chart = os.path.join(GRAFICOS_DIR, "pagespeed-comparativo.png")
    add_image_safe(slide, chart, Inches(0.5), Inches(1.6), Inches(9), Inches(5.2))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 4: CORE WEB VITALS
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Core Web Vitals", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Métricas clave de experiencia de usuario - Antes vs Después", font_size=14, color=GRAY)

    cwv_chart = os.path.join(GRAFICOS_DIR, "core-web-vitals.png")
    add_image_safe(slide, cwv_chart, Inches(0.5), Inches(1.6), Inches(9), Inches(5.2))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 5: SEGURIDAD Y ESTABILIDAD
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Seguridad y Estabilidad", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Protegiendo la información y asegurando el funcionamiento", font_size=14, color=GRAY)

    items = [
        "Clave secreta protegida fuera del código fuente",
        "Correo funcional: mensajes llegan a info@magnaingenieriaytopografia.com",
        "Dependencias limpias y actualizadas (19 librerías actualizadas)",
        "56 pruebas automáticas que verifican todo el sitio",
        "Certificado SSL (HTTPS) - candado verde en el navegador",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(9), Inches(4.5), items, font_size=16, color=DARK)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 6: INFRAESTRUCTURA
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Infraestructura Moderna", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Los cimientos del sitio web ahora son profesionales", font_size=14, color=GRAY)

    items = [
        "Docker: despliegue con un solo comando, entornos idénticos",
        "PostgreSQL 16: base de datos profesional (la misma que usa Instagram)",
        "Nginx: proxy inverso que acelera y protege el sitio",
        "Caché: archivos estáticos guardados por hasta 1 año",
        "HTTP/2: múltiples transferencias simultáneas",
        "Compresión Gzip: archivos más pequeños, carga más rápida",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(9), Inches(4.5), items, font_size=15, color=DARK)

    infra_img = os.path.join(GRAFICOS_DIR, "legacy-pagespeed-1.png")
    add_image_safe(slide, infra_img, Inches(6), Inches(1.8), Inches(3.5), Inches(2.5))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 7: HERO SLIDER
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Nuevo Hero Slider", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Slides independientes para promocionar servicios clave", font_size=14, color=GRAY)

    hero_img = os.path.join(SCREENSHOTS_DIR, "home-hero.png")
    add_image_safe(slide, hero_img, Inches(0.5), Inches(1.6), Inches(9), Inches(4.5))
    add_text_box(slide, Inches(0.5), Inches(6.3), Inches(9), Inches(0.5),
                 "✓ Responsive (desktop, tablet, móvil)  ✓ Imágenes adaptativas  ✓ Navegación accesible",
                 font_size=12, color=GRAY, alignment=PP_ALIGN.CENTER)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 8: SERVICIOS
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Servicios y Subservicios", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Sección de servicios con tarjetas animadas y subservicios detallados", font_size=14, color=GRAY)

    items = [
        "Tarjetas de servicio con animación escalonada (Framer Motion)",
        "Imágenes con lazy loading para carga rápida",
        "Subservicios con imágenes adaptativas (desktop, tablet, celular)",
        "Generación automática de variantes de imagen",
        "Slider de subservicios con Swiper",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(4.5), Inches(3), items, font_size=14, color=DARK)

    serv_img = os.path.join(SCREENSHOTS_DIR, "home-servicios.png")
    add_image_safe(slide, serv_img, Inches(5.2), Inches(1.6), Inches(4.5), Inches(4))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 9: REDISEÑO CONTACTO
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Rediseño de Contacto", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Formulario funcional + tarjetas informativas + mapa interactivo", font_size=14, color=GRAY)

    items = [
        "4 tarjetas interactivas con iconos (Dirección, Teléfono, Email, Horarios)",
        "Formulario que envía correos a info@magnaingenieriaytopografia.com",
        "Casilla de aceptación de política de datos",
        "Mapa Leaflet con coordenadas configurables",
        "Notificación visual al enviar el formulario",
        "SEO: Open Graph y Twitter Cards",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(9), Inches(3.5), items, font_size=15, color=DARK)

    contacto_img = os.path.join(SCREENSHOTS_DIR, "contacto.png")
    add_image_safe(slide, contacto_img, Inches(0.5), Inches(4.5), Inches(9), Inches(2.5))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 10: INFORMACIÓN CORPORATIVA
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Información Corporativa", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Misión, Visión, Valores, Equipo y Herramienta Logística", font_size=14, color=GRAY)

    items = [
        "Misión, Visión y Valores institucionales actualizados",
        "Nuevo bloque: Herramienta logística y personal certificado",
        "Miembros del equipo con tarjetas interactivas",
        "Dirección y teléfonos actualizados",
        "Sección Quiénes Somos completamente renovada",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(4.5), Inches(3), items, font_size=15, color=DARK)

    qs_img = os.path.join(SCREENSHOTS_DIR, "home-quienes-somos.png")
    add_image_safe(slide, qs_img, Inches(5.2), Inches(1.6), Inches(4.5), Inches(3.5))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 11: BLOG
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Blog", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Corrección de errores, mejoras visuales, SEO y accesibilidad", font_size=14, color=GRAY)

    items = [
        "5 errores críticos corregidos (imágenes, búsqueda, formato)",
        "Tarjetas con diseño destacado para artículos importantes",
        "SEO: Open Graph, Twitter Cards, Schema.org, Sitemap",
        "Accesibilidad: aria-labels, contraste, teclado",
        "Búsqueda con debounce y categorías",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(4.5), Inches(3), items, font_size=15, color=DARK)

    blog_img = os.path.join(SCREENSHOTS_DIR, "blog.png")
    add_image_safe(slide, blog_img, Inches(5.2), Inches(1.6), Inches(4.5), Inches(3.5))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 12: TIENDA
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Tienda Online - Magnatienda", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Remodelación visual completa en 12 fases", font_size=14, color=GRAY)

    items = [
        "Tarjetas de producto con hover, sombra y borde dorado",
        "Precios en COP, calificación con estrellas",
        "Carrito responsivo con resumen pegajoso",
        "Checkout con PayPal integrado",
        "Navegación con categorías y búsqueda",
        "Perfil de usuario e historial de pedidos",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(4.5), Inches(3), items, font_size=14, color=DARK)

    tienda_img = os.path.join(SCREENSHOTS_DIR, "tienda-home.png")
    add_image_safe(slide, tienda_img, Inches(5.2), Inches(1.6), Inches(4.5), Inches(3.5))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 13: MÁS MEJORAS
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Otras Mejoras", font_size=32, bold=True, color=NAVY)

    items = [
        "ProgressiveBackground: imágenes borrosas que se vuelven nítidas",
        "Estadísticas animadas (clientes, proyectos, cobertura nacional)",
        "Logos de clientes en carrusel con paralaje",
        "Preguntas frecuentes con acordeón interactivo",
        "Política de datos personales (Ley 1581 de 2012)",
        "Footer completo con contacto, redes sociales y mapa del sitio",
        "Barra de navegación sticky con diseño responsive",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.3), Inches(9), Inches(5), items, font_size=16, color=DARK)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 14: REDUCCIÓN CSS
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Optimización de Rendimiento", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "CSS reducido en 56%, JavaScript dividido en 15 bloques", font_size=14, color=GRAY)

    css_chart = os.path.join(GRAFICOS_DIR, "reduccion-css.png")
    add_image_safe(slide, css_chart, Inches(0.5), Inches(1.6), Inches(5), Inches(4.5))

    items = [
        "PurgeCSS: 232 KB → 102.5 KB (56% menos)",
        "Code Splitting: ~15 bloques bajo demanda",
        "Lazy Loading en imágenes y secciones",
        "Sourcemaps desactivados (-50% JS)",
        "Preconnect + DNS-prefetch para Google Fonts",
    ]
    add_multi_text(slide, Inches(5.8), Inches(1.8), Inches(4), Inches(4), items, font_size=14, color=DARK)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 15: PROYECTOS
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Proyectos", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "Galería de proyectos con filtros por categoría", font_size=14, color=GRAY)

    items = [
        "Tarjetas de proyecto con filtro por tipo",
        "Imágenes con lazy loading y srcset responsive",
        "Animaciones de entrada con Framer Motion",
        "Carrusel de imágenes en detalle del proyecto",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(4.5), Inches(2.5), items, font_size=15, color=DARK)

    proy_img = os.path.join(SCREENSHOTS_DIR, "proyectos.png")
    add_image_safe(slide, proy_img, Inches(5.2), Inches(1.6), Inches(4.5), Inches(4))
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 16: BENEFICIOS NEGOCIO
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    slide_num += 1

    add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.8),
                 "Beneficios para tu Negocio", font_size=32, bold=True, color=NAVY)
    add_text_box(slide, Inches(0.5), Inches(1.0), Inches(9), Inches(0.5),
                 "¿Qué significa todo esto para MAGNA?", font_size=14, color=GRAY)

    items = [
        "Más clientes desde Google (mejor SEO y velocidad)",
        "Profesionalismo y confianza (diseño moderno + HTTPS)",
        "Más consultas recibidas (formulario funcional)",
        "Menos llamadas innecesarias (FAQ + blog informativo)",
        "Nuevo canal de ventas (tienda online con PayPal)",
        "Sitio siempre disponible (infraestructura profesional)",
        "Protección legal (política de datos personales)",
        "Actualización sencilla (panel de administración)",
    ]
    add_multi_text(slide, Inches(0.5), Inches(1.8), Inches(9), Inches(4.5), items, font_size=15, color=DARK)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # SLIDE 17: CIERRE
    # ============================================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    slide_num += 1

    add_text_box(slide, Inches(1), Inches(2), Inches(8), Inches(1),
                 "¡Gracias!", font_size=48, bold=True, color=GOLD, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(3.3), Inches(8), Inches(1.5),
                 "MAGNA Ingeniería y Topografía\n\n"
                 "info@magnaingenieriaytopografia.com\n"
                 "magnaingenieriaytopografia.com", font_size=18, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_slide_number(slide, slide_num, TOTAL_SLIDES)

    # ============================================================
    # Save
    # ============================================================
    prs.save(OUTPUT_PATH)
    print(f"Presentación PowerPoint guardada: {OUTPUT_PATH}")

if __name__ == "__main__":
    create_presentation()
